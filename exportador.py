"""
Exportador Tableau Server / Tableau Cloud
=========================================

Descarga a una carpeta local el contenido de uno o más sitios de Tableau:
  - Workbooks (.twbx / .twb)
  - Fuentes de datos publicadas (.tdsx / .tds)
  - Flujos de Prep (.tflx / .tfl)

Flujo:
  Fase 1 - Inventario: lista todo el contenido y genera inventario_<fecha>.xlsx antes de descargar
           (proyectos, workbooks, vistas, fuentes, conexiones, flujos, usuarios, grupos, tareas,
           suscripciones y campos por hoja como en el proyecto Copec).
  Fase 2 - Descarga: baja los archivos y actualiza en el inventario el estado de cada elemento.

Salidas:
  - inventario_<fecha>.xlsx : inventario completo con la paleta de colores Backspace
  - pendientes_<fecha>.xlsx : todo lo que NO se pudo descargar, con su URL y la acción sugerida
                              para hacerlo manualmente
  - logs/ejecucion_<fecha>.log

Uso:
    python exportador.py                    (usa config.ini)
    python exportador.py --config otro.ini
    python exportador.py --solo-inventario  (no descarga archivos)
    python exportador.py --prueba           (modo prueba: solo 5 elementos por tipo)
"""

import argparse
import configparser
import getpass
import logging
import os
import re
import sys
import time
from datetime import date, datetime

import pandas as pd
import tableauserverclient as TSC
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

# --------------------------- #
#        CONFIGURACIÓN        #
# --------------------------- #

EXTENSIONES = {
    "Workbook": [".twbx", ".twb"],
    "Fuente de datos": [".tdsx", ".tds"],
    "Flujo": [".tflx", ".tfl"],
}

CARPETAS = {
    "Workbook": "Workbooks",
    "Fuente de datos": "Fuentes de datos",
    "Flujo": "Flujos",
}

# Query de campos por hoja (misma estructura que el extractor de Copec, filtrando por LUID)
GRAPHQL_CAMPOS_WORKBOOK = """
query CamposPorWorkbook($luid: String) {
  workbooks(filter: {luid: $luid}) {
    name
    luid
    sheets {
      id
      name
      sheetFieldInstances {
        name
        __typename
        ... on ColumnField { dataType }
        ... on CalculatedField { dataType }
        ... on GroupField { dataType }
        ... on DatasourceField {
          remoteField {
            ... on ColumnField { dataType }
          }
        }
      }
    }
  }
}
"""

log = logging.getLogger("exportador")
pendientes = []      # lo que falta descargar / revisar manualmente


# ---------------------- #
#       UTILIDADES       #
# ---------------------- #

def leer_config(ruta):
    if not os.path.exists(ruta):
        sys.exit(f"No se encontró el archivo de configuración '{ruta}'")
    cfg = configparser.ConfigParser()
    cfg.read(ruta, encoding="utf-8")
    return cfg


def configurar_logs(carpeta, marca):
    os.makedirs(os.path.join(carpeta, "logs"), exist_ok=True)
    ruta_log = os.path.join(carpeta, "logs", f"ejecucion_{marca}.log")
    formato = logging.Formatter("%(asctime)s | %(levelname)-7s | %(message)s", "%Y-%m-%d %H:%M:%S")

    # La consola de Windows (cp1252) no soporta emojis/tildes en todos los casos
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    archivo = logging.FileHandler(ruta_log, encoding="utf-8")
    archivo.setFormatter(formato)
    consola = logging.StreamHandler(sys.stdout)
    consola.setFormatter(logging.Formatter("%(message)s"))

    log.setLevel(logging.INFO)
    log.addHandler(archivo)
    log.addHandler(consola)
    return ruta_log


def limpiar_nombre(nombre, largo_max=80):
    """Deja un nombre válido para carpeta/archivo en Windows."""
    nombre = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(nombre or "sin_nombre")).strip(" .")
    return (nombre or "sin_nombre")[:largo_max]


def ruta_windows(ruta):
    """Permite rutas de más de 260 caracteres en Windows."""
    ruta = os.path.abspath(ruta)
    if os.name == "nt" and not ruta.startswith("\\\\?\\"):
        return "\\\\?\\" + ruta
    return ruta


def fecha(valor):
    """Fechas sin zona horaria (Excel no las soporta)."""
    return valor.replace(tzinfo=None) if isinstance(valor, datetime) else valor


def listar(endpoint, **kwargs):
    """Lista todos los elementos de un endpoint paginado."""
    return list(TSC.Pager(endpoint, **kwargs))


def intentar(descripcion, sitio, funcion, por_defecto=None):
    """Ejecuta una consulta del inventario; si falla, la deja registrada y sigue."""
    try:
        return funcion()
    except Exception as e:
        log.warning(f"   ⚠ No se pudo obtener {descripcion}: {texto_error(e)}")
        registrar_pendiente(sitio, "Inventario", "", descripcion, "", "", e,
                            "Revisar permisos del token (se requiere rol Administrador de sitio/servidor).")
        return por_defecto


def accion_sugerida(error, tipo):
    texto = str(error).lower()
    if "403" in texto or "forbidden" in texto or "permission" in texto:
        return f"Sin permiso para descargar el {tipo.lower()}. Descargarlo con un administrador o dar permiso de descarga al usuario del token."
    if "404" in texto or "not found" in texto:
        return "El contenido no existe o no es visible para el usuario del token. Verificar en la web."
    if "timed out" in texto or "timeout" in texto or "connection" in texto:
        return "Error de red / tiempo de espera (archivo muy pesado). Re-ejecutar el exportador o descargar manualmente desde la URL."
    if "too long" in texto or "errno 2" in texto or "errno 22" in texto:
        return "Problema con la ruta local del archivo. Descargar manualmente desde la URL."
    return "Descargar manualmente desde la URL web."


def texto_error(error):
    """Mensaje de error en una sola línea (los de Tableau vienen con saltos de línea)."""
    return " ".join(str(error).split())


def registrar_pendiente(sitio, tipo, proyecto, nombre, item_id, url, error, accion=None):
    pendientes.append({
        "fecha": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "sitio": sitio,
        "tipo": tipo,
        "proyecto": proyecto,
        "nombre": nombre,
        "id": item_id,
        "url_web": url,
        "error": texto_error(error)[:500],
        "accion_sugerida": accion or accion_sugerida(error, tipo),
    })


# ---------------------- #
#        CONEXIÓN        #
# ---------------------- #

class Conexion:
    """Maneja el login con PAT y la reconexión si la sesión expira durante descargas largas."""

    def __init__(self, url, token_name, token_value, verify_ssl):
        self.url = url.rstrip("/")
        self.token_name = token_name
        self.token_value = token_value
        self.server = TSC.Server(self.url, use_server_version=False)
        if not verify_ssl:
            import urllib3
            urllib3.disable_warnings()
            self.server.add_http_options({"verify": False})
        self.server.use_server_version()
        self.sitio = None

    def conectar(self, sitio):
        if self.server.is_signed_in():
            self.server.auth.sign_out()
        auth = TSC.PersonalAccessTokenAuth(self.token_name, self.token_value, site_id=sitio)
        self.server.auth.sign_in(auth)
        self.sitio = sitio

    def reconectar(self):
        log.info("   🔄 Renovando sesión...")
        self.conectar(self.sitio)

    def desconectar(self):
        if self.server.is_signed_in():
            self.server.auth.sign_out()

    def url_vista(self, content_url):
        # content_url de vista: "Workbook/sheets/Vista" -> #/views/Workbook/Vista
        ruta = (content_url or "").replace("/sheets/", "/")
        prefijo = f"site/{self.sitio}/" if self.sitio else ""
        return f"{self.url}/#/{prefijo}views/{ruta}"


# ---------------------- #
#       DESCARGAS        #
# ---------------------- #

def descargar(conexion, tipo, item, carpeta_tipo, ruta_proyecto, opciones):
    """Descarga un elemento. Retorna (estado, ruta_local)."""
    carpeta = os.path.join(carpeta_tipo, *[limpiar_nombre(p) for p in ruta_proyecto.split("/")])
    base = os.path.join(carpeta, limpiar_nombre(item.name))

    # Si ya existe (de una ejecución anterior) se salta
    for ext in EXTENSIONES[tipo]:
        if os.path.exists(ruta_windows(base + ext)) and not opciones["sobrescribir"]:
            return "Ya existía", base + ext

    os.makedirs(ruta_windows(carpeta), exist_ok=True)
    temporal = ruta_windows(base + "__descargando")
    server = conexion.server

    for intento in range(2):
        try:
            if tipo == "Workbook":
                ruta = server.workbooks.download(item.id, filepath=temporal, include_extract=opciones["incluir_extractos"])
            elif tipo == "Fuente de datos":
                ruta = server.datasources.download(item.id, filepath=temporal, include_extract=opciones["incluir_extractos"])
            else:
                ruta = server.flows.download(item.id, filepath=temporal)

            extension = os.path.splitext(ruta)[1]
            final = ruta_windows(base + extension)
            os.replace(ruta, final)
            return "Descargado", base + extension

        except TSC.ServerResponseError as e:
            borrar_temporales(temporal, tipo)
            if str(e.code).startswith("401") and intento == 0:
                conexion.reconectar()
                continue
            raise
        except Exception:
            borrar_temporales(temporal, tipo)
            if intento == 0:
                time.sleep(3)
                continue
            raise


def borrar_temporales(temporal, tipo):
    for ext in EXTENSIONES[tipo]:
        if os.path.exists(temporal + ext):
            os.remove(temporal + ext)


def descargar_lista(conexion, tipo, items, carpeta_sitio, rutas_items, opciones):
    """Descarga una lista de workbooks/fuentes/flujos y retorna {id: (estado, archivo)}."""
    resultados = {}
    carpeta_tipo = os.path.join(carpeta_sitio, CARPETAS[tipo])
    total = len(items)

    for i, item in enumerate(items, 1):
        proyecto = rutas_items[item.id]
        inicio = time.time()
        try:
            estado, archivo = descargar(conexion, tipo, item, carpeta_tipo, proyecto, opciones)
            tamano = os.path.getsize(ruta_windows(archivo)) / 1024 / 1024
            log.info(f"   [{i}/{total}] {tipo} '{proyecto}/{item.name}' -> {estado} "
                     f"({tamano:.1f} MB, {time.time() - inicio:.0f}s)")
        except Exception as e:
            estado, archivo = "Error", ""
            log.error(f"   [{i}/{total}] {tipo} '{proyecto}/{item.name}' -> ERROR: {texto_error(e)}")
            registrar_pendiente(conexion.sitio, tipo, proyecto, item.name, item.id,
                                getattr(item, "webpage_url", ""), e)
        resultados[item.id] = (estado, archivo)

    return resultados


def seleccionar_prueba(items, rutas_items, cantidad):
    """Modo prueba: toma los N elementos más livianos, priorizando los que no están en espacios personales."""
    def orden(item):
        personal = rutas_items[item.id].startswith("Espacio personal")
        return personal, getattr(item, "size", None) or 0
    return sorted(items, key=orden)[:cantidad]


def descargar_sitio(conexion, sitio_info, opciones):
    """Fase 2: descarga los archivos de un sitio ya inventariado."""
    contenido, rutas_items = sitio_info["contenido"], sitio_info["rutas_items"]
    descargas = {}
    for tipo, opcion in (("Workbook", "descargar_workbooks"), ("Fuente de datos", "descargar_fuentes"),
                         ("Flujo", "descargar_flujos")):
        items = sitio_info["seleccion"][tipo]
        if not opciones[opcion] or not items:
            continue
        log.info(f"⬇  Descargando {len(items)} de {len(contenido[tipo])} {CARPETAS[tipo].lower()}...")
        descargas.update(descargar_lista(conexion, tipo, items, sitio_info["carpeta"], rutas_items, opciones))
    return descargas


# ---------------------- #
#       INVENTARIO       #
# ---------------------- #

def rutas_proyectos(proyectos):
    """Arma la ruta completa de cada proyecto: 'Padre/Hijo/Nieto'."""
    por_id = {p.id: p for p in proyectos}
    rutas = {}
    for p in proyectos:
        partes, actual = [], p
        while actual and len(partes) < 50:
            partes.insert(0, actual.name)
            actual = por_id.get(actual.parent_id)
        rutas[p.id] = "/".join(partes)
    return rutas


def filtrar_por_proyecto(items, rutas, filtro):
    if not filtro:
        return items
    return [i for i in items if rutas.get(i.project_id, "").split("/")[0] in filtro]


def campos_por_hoja(conexion, workbooks, rutas_items):
    """Inventario de campos usados en cada hoja (mismo formato que el proyecto Copec)."""
    registros = []
    errores_seguidos = 0
    total = len(workbooks)
    for i, wb in enumerate(workbooks, 1):
        log.info(f"   [{i}/{total}] Campos de '{wb.name}'")
        try:
            try:
                respuesta = conexion.server.metadata.query(GRAPHQL_CAMPOS_WORKBOOK, variables={"luid": wb.id})
            except TSC.ServerResponseError as e:
                if not str(e.code).startswith("401"):
                    raise
                conexion.reconectar()
                respuesta = conexion.server.metadata.query(GRAPHQL_CAMPOS_WORKBOOK, variables={"luid": wb.id})

            if respuesta.get("errors"):
                raise RuntimeError(respuesta["errors"])

            for wb_meta in respuesta.get("data", {}).get("workbooks", []):
                for hoja in wb_meta.get("sheets", []):
                    for campo in hoja.get("sheetFieldInstances", []):
                        tipo = campo.get("__typename")
                        if tipo == "DatasourceField":
                            data_type = (campo.get("remoteField") or {}).get("dataType")
                        else:
                            data_type = campo.get("dataType")
                        registros.append({
                            "site_name": conexion.sitio or "Default",
                            "project": rutas_items[wb.id],
                            "workbook_name": wb_meta.get("name"),
                            "workbook_id": wb_meta.get("luid"),
                            "sheet_name": hoja.get("name"),
                            "sheet_id": hoja.get("id"),
                            "field_name": campo.get("name"),
                            "field_type": tipo,
                            "data_type": data_type,
                            "extraction_date": date.today(),
                        })
            errores_seguidos = 0
        except Exception as e:
            errores_seguidos += 1
            log.warning(f"   ⚠ Error leyendo campos de '{wb.name}': {texto_error(e)}")
            registrar_pendiente(conexion.sitio, "Inventario campos", rutas_items[wb.id], wb.name, wb.id,
                                wb.webpage_url, e,
                                "Verificar que la Metadata API esté habilitada "
                                "(tsm maintenance metadata-services enable) y que el token vea el workbook.")
            if errores_seguidos >= 5 and not registros:
                log.warning("   ⚠ La Metadata API no responde; se omite el inventario de campos de este sitio.")
                break
    return registros


def leer_inventario(conexion, opciones):
    """Fase 1: lista todo el contenido del sitio y arma las tablas del inventario (sin descargar nada)."""
    server = conexion.server
    sitio = conexion.sitio
    nombre_sitio = sitio or "Default"
    filtro = opciones["filtro_proyectos"]
    prueba = opciones["modo_prueba"]
    tablas = {}

    log.info("📂 Leyendo proyectos, usuarios y grupos...")
    proyectos = listar(server.projects)
    rutas = rutas_proyectos(proyectos)
    usuarios = intentar("usuarios", sitio, lambda: listar(server.users), [])
    nombre_usuario = {u.id: u.name for u in usuarios}

    tablas["Proyectos"] = [{
        "sitio": nombre_sitio, "id": p.id, "nombre": p.name, "ruta": rutas[p.id],
        "proyecto_padre": rutas.get(p.parent_id, ""), "descripcion": p.description,
        "modo_permisos": p.content_permissions, "propietario": nombre_usuario.get(getattr(p, "owner_id", None), ""),
    } for p in proyectos if not filtro or rutas[p.id].split("/")[0] in filtro]

    tablas["Usuarios"] = [{
        "sitio": nombre_sitio, "id": u.id, "usuario": u.name, "nombre_completo": u.fullname, "email": u.email,
        "rol_sitio": u.site_role, "autenticacion": u.auth_setting, "ultimo_login": fecha(u.last_login),
    } for u in usuarios]

    grupos = intentar("grupos", sitio, lambda: listar(server.groups), [])
    tablas["Grupos"], tablas["Usuarios por grupo"] = [], []
    for g in grupos:
        miembros = intentar(f"miembros del grupo {g.name}", sitio, lambda: _miembros_grupo(server, g), [])
        tablas["Grupos"].append({
            "sitio": nombre_sitio, "id": g.id, "nombre": g.name, "dominio": g.domain_name,
            "rol_minimo": g.minimum_site_role, "cantidad_usuarios": len(miembros),
        })
        tablas["Usuarios por grupo"] += [{"sitio": nombre_sitio, "grupo": g.name, "usuario": u.name} for u in miembros]

    # --- Workbooks, vistas, fuentes de datos y flujos ---
    log.info("📚 Leyendo workbooks, vistas, fuentes de datos y flujos...")
    workbooks = filtrar_por_proyecto(listar(server.workbooks), rutas, filtro)
    fuentes = filtrar_por_proyecto(listar(server.datasources), rutas, filtro)
    flujos = filtrar_por_proyecto(intentar("flujos", sitio, lambda: listar(server.flows), []), rutas, filtro)
    vistas = intentar("vistas con uso", sitio, lambda: listar(server.views, usage=True), None)
    if vistas is None:
        vistas = intentar("vistas", sitio, lambda: listar(server.views), [])
    ids_wb = {w.id: w for w in workbooks}
    vistas = [v for v in vistas if v.workbook_id in ids_wb]
    log.info(f"   {len(proyectos)} proyectos | {len(workbooks)} workbooks | {len(vistas)} vistas | "
             f"{len(fuentes)} fuentes | {len(flujos)} flujos | {len(usuarios)} usuarios")

    # Ruta de cada elemento. Los espacios personales se separan por propietario para no mezclar archivos.
    rutas_items = {}
    for item in workbooks + fuentes + flujos:
        rutas_items[item.id] = rutas.get(item.project_id) or \
            f"Espacio personal/{nombre_usuario.get(item.owner_id, item.owner_id)}"

    contenido = {"Workbook": workbooks, "Fuente de datos": fuentes, "Flujo": flujos}
    if prueba:
        seleccion = {tipo: seleccionar_prueba(items, rutas_items, opciones["cantidad_prueba"])
                     for tipo, items in contenido.items()}
        log.info(f"🧪 Modo prueba: se procesan {opciones['cantidad_prueba']} elementos por tipo")
    else:
        seleccion = contenido
    ids_seleccion = {i.id for items in seleccion.values() for i in items}
    estado_inicial = "Pendiente de descarga"

    def estado(item_id):
        return estado_inicial if item_id in ids_seleccion else "No incluido (modo prueba)"

    # --- Tablas de contenido ---
    log.info("🔌 Leyendo conexiones de workbooks y fuentes...")
    tablas["Workbooks"], tablas["Fuentes de datos"], tablas["Conexiones"] = [], [], []
    for wb in workbooks:
        tablas["Workbooks"].append({
            "sitio": nombre_sitio, "id": wb.id, "nombre": wb.name, "proyecto": rutas_items[wb.id],
            "propietario": nombre_usuario.get(wb.owner_id, wb.owner_id), "creado": fecha(wb.created_at),
            "actualizado": fecha(wb.updated_at), "tamano": wb.size, "mostrar_pestanas": wb.show_tabs,
            "descripcion": wb.description, "etiquetas": ", ".join(sorted(wb.tags or [])),
            "url_web": wb.webpage_url, "estado_descarga": estado(wb.id), "archivo_local": "",
        })
        if wb.id in ids_seleccion:
            conexiones = intentar(f"conexiones del workbook {wb.name}", sitio,
                                  lambda: server.workbooks.populate_connections(wb) or wb.connections, [])
            tablas["Conexiones"] += filas_conexiones(nombre_sitio, "Workbook", wb, rutas_items, conexiones)

    for ds in fuentes:
        tablas["Fuentes de datos"].append({
            "sitio": nombre_sitio, "id": ds.id, "nombre": ds.name, "proyecto": rutas_items[ds.id],
            "propietario": nombre_usuario.get(ds.owner_id, ds.owner_id), "tipo": ds.datasource_type,
            "tiene_extracto": ds.has_extracts, "certificada": ds.certified, "creado": fecha(ds.created_at),
            "actualizado": fecha(ds.updated_at), "descripcion": ds.description,
            "etiquetas": ", ".join(sorted(ds.tags or [])), "url_web": ds.webpage_url,
            "estado_descarga": estado(ds.id), "archivo_local": "",
        })
        if ds.id in ids_seleccion:
            conexiones = intentar(f"conexiones de la fuente {ds.name}", sitio,
                                  lambda: server.datasources.populate_connections(ds) or ds.connections, [])
            tablas["Conexiones"] += filas_conexiones(nombre_sitio, "Fuente de datos", ds, rutas_items, conexiones)

    tablas["Vistas"] = [{
        "sitio": nombre_sitio, "id": v.id, "nombre": v.name, "workbook": ids_wb[v.workbook_id].name,
        "proyecto": rutas_items[v.workbook_id],
        "propietario": nombre_usuario.get(v.owner_id, v.owner_id), "tipo_hoja": getattr(v, "sheet_type", None),
        "total_visitas": getattr(v, "_total_views", None),
        "creado": fecha(v.created_at), "actualizado": fecha(v.updated_at), "url_web": conexion.url_vista(v.content_url),
    } for v in vistas]

    tablas["Flujos"] = [{
        "sitio": nombre_sitio, "id": f.id, "nombre": f.name, "proyecto": rutas_items[f.id],
        "propietario": nombre_usuario.get(f.owner_id, f.owner_id), "creado": fecha(f.created_at),
        "actualizado": fecha(f.updated_at), "descripcion": f.description, "url_web": f.webpage_url,
        "estado_descarga": estado(f.id), "archivo_local": "",
    } for f in flujos]

    # --- Tareas y suscripciones ---
    nombres = {i.id: i.name for i in workbooks + fuentes + flujos}
    tareas = intentar("tareas de actualización de extractos", sitio, lambda: list(server.tasks.get()[0]), [])
    tablas["Tareas extraccion"] = [{
        "sitio": nombre_sitio, "id": t.id, "tipo_tarea": t.task_type, "tipo_contenido": t.target.type,
        "contenido": nombres.get(t.target.id, t.target.id), "prioridad": t.priority,
        "fallos_consecutivos": t.consecutive_failed_count,
        "programacion": getattr(t.schedule_item, "name", None) if t.schedule_item else None,
        "frecuencia": describir_programacion(t.schedule_item),
        "proxima_ejecucion": fecha(getattr(t.schedule_item, "next_run_at", None)) if t.schedule_item else None,
    } for t in tareas]

    suscripciones = intentar("suscripciones", sitio, lambda: listar(server.subscriptions), [])
    tablas["Suscripciones"] = [{
        "sitio": nombre_sitio, "id": s.id, "asunto": s.subject,
        "usuario": nombre_usuario.get(s.user_id, s.user_id), "tipo_contenido": s.target.type,
        "contenido": nombres.get(s.target.id, s.target.id), "programacion": getattr(s, "schedule_id", None),
    } for s in suscripciones]

    # --- Campos por hoja (inventario Copec) ---
    tablas["Campos"] = []
    if opciones["inventario_campos"] and seleccion["Workbook"]:
        log.info("🧬 Leyendo campos por hoja desde la Metadata API...")
        tablas["Campos"] = campos_por_hoja(conexion, seleccion["Workbook"], rutas_items)

    tablas["_resumen"] = {
        "sitio": nombre_sitio, "proyectos": len(tablas["Proyectos"]), "workbooks": len(workbooks),
        "vistas": len(vistas), "fuentes_de_datos": len(fuentes), "flujos": len(flujos),
        "usuarios": len(usuarios), "grupos": len(grupos),
        "descargados": 0, "ya_existian": 0, "con_error": 0,
    }
    return {"sitio": sitio, "tablas": tablas, "contenido": contenido, "seleccion": seleccion,
            "rutas_items": rutas_items}


def aplicar_descargas(tablas, descargas, solo_inventario):
    """Actualiza en el inventario el estado de descarga de cada elemento."""
    for hoja in ("Workbooks", "Fuentes de datos", "Flujos"):
        for fila in tablas[hoja]:
            if fila["id"] in descargas:
                fila["estado_descarga"], fila["archivo_local"] = descargas[fila["id"]]
            elif solo_inventario and fila["estado_descarga"] == "Pendiente de descarga":
                fila["estado_descarga"] = "No descargado (solo inventario)"
    estados = [e for e, _ in descargas.values()]
    tablas["_resumen"].update(descargados=estados.count("Descargado"), ya_existian=estados.count("Ya existía"),
                              con_error=estados.count("Error"))


def _miembros_grupo(server, grupo):
    server.groups.populate_users(grupo)
    return list(grupo.users)


def filas_conexiones(sitio, tipo, item, rutas_items, conexiones):
    return [{
        "sitio": sitio, "tipo_contenido": tipo, "contenido": item.name,
        "proyecto": rutas_items[item.id], "id_conexion": c.id,
        "tipo_conexion": c.connection_type, "servidor": c.server_address, "puerto": c.server_port,
        "usuario_bd": c.username, "password_embebida": c.embed_password,
        "fuente_datos": getattr(c, "datasource_name", None),
    } for c in (conexiones or [])]


def describir_programacion(programacion):
    if not programacion:
        return None
    intervalo = getattr(programacion, "interval_item", None)
    if intervalo is None:
        return None
    partes = [type(intervalo).__name__.replace("Interval", "")]
    for atributo in ("start_time", "interval", "_interval"):
        valor = getattr(intervalo, atributo, None)
        if valor:
            partes.append(str(valor))
    return " ".join(partes)


# ---------------------- #
#     SALIDAS (EXCEL)    #
# ---------------------- #

# Paleta Backspace (tomada de backspace.cl)
NEGRO = "313131"
VERDE = "99CC33"
LIMA = "AEEE2F"
GRIS = "838080"
CLARO = "EEEEEE"
ROJO = "EF4444"

COLOR_ESTADO = {    # relleno, color de letra
    "Descargado": ("F0FAE0", "4D6B12"),
    "Ya existía": ("F0FAE0", "4D6B12"),
    "Error": ("FEE2E2", "991B1B"),
    "Pendiente de descarga": ("FFFBEB", "8A6D1A"),
}

HOJAS_CONTENIDO = ["Proyectos", "Workbooks", "Vistas", "Fuentes de datos", "Conexiones", "Flujos", "Campos"]
HOJAS_ADMIN = ["Usuarios", "Grupos", "Usuarios por grupo", "Tareas extraccion", "Suscripciones"]


def guardar_excel(ruta, funcion):
    """Guarda un Excel; si el archivo está abierto (típico en reuniones) lo guarda con otro nombre."""
    try:
        funcion(ruta)
        return ruta
    except PermissionError:
        alternativa = ruta.replace(".xlsx", f"_{datetime.now():%H%M%S}.xlsx")
        log.warning(f"   ⚠ '{os.path.basename(ruta)}' está abierto en Excel; se guarda como '{os.path.basename(alternativa)}'")
        funcion(alternativa)
        return alternativa


def guardar_inventario(sitios_info, info, ruta):
    def escribir(destino):
        with pd.ExcelWriter(destino, engine="openpyxl") as writer:
            resumen = pd.DataFrame([s["tablas"]["_resumen"] for s in sitios_info])
            escribir_resumen(writer, "Inventario Tableau", info, resumen)
            for hoja in HOJAS_CONTENIDO + HOJAS_ADMIN:
                filas = [f for s in sitios_info for f in s["tablas"].get(hoja, [])]
                escribir_hoja(writer, hoja, pd.DataFrame(filas), VERDE if hoja in HOJAS_CONTENIDO else GRIS)
            escribir_hoja(writer, "Pendientes", tabla_pendientes(), ROJO)
    return guardar_excel(ruta, escribir)


def guardar_pendientes(info, ruta):
    def escribir(destino):
        df = tabla_pendientes()
        with pd.ExcelWriter(destino, engine="openpyxl") as writer:
            if pendientes:
                resumen = (df.groupby(["sitio", "tipo", "accion_sugerida"]).size()
                           .reset_index(name="cantidad").sort_values("cantidad", ascending=False))
            else:
                resumen = pd.DataFrame({"resultado": ["No hay elementos pendientes 🎉"]})
            escribir_resumen(writer, "Pendientes de descarga manual", info, resumen)
            escribir_hoja(writer, "Pendientes", df, ROJO)
    return guardar_excel(ruta, escribir)


def tabla_pendientes():
    columnas = ["fecha", "sitio", "tipo", "proyecto", "nombre", "id", "url_web", "error", "accion_sugerida"]
    return pd.DataFrame(pendientes, columns=columnas)


def escribir_resumen(writer, titulo, info, tabla):
    """Hoja de portada: título con la paleta Backspace, datos de la ejecución y tabla resumen."""
    from openpyxl.styles import Alignment, Font, PatternFill

    fila_tabla = len(info) + 5
    limpiar_df(tabla).to_excel(writer, sheet_name="Resumen", index=False, startrow=fila_tabla - 1)
    ws = writer.sheets["Resumen"]
    ws.sheet_properties.tabColor = NEGRO
    ws.sheet_view.showGridLines = False

    ancho = max(len(tabla.columns), 2)
    for col in range(1, ancho + 1):
        ws.cell(row=1, column=col).fill = PatternFill("solid", fgColor=NEGRO)
        ws.cell(row=2, column=col).fill = PatternFill("solid", fgColor=VERDE)
    ws.cell(row=1, column=1, value=titulo).font = Font(size=16, bold=True, color=LIMA)
    ws.cell(row=1, column=1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32
    ws.row_dimensions[2].height = 4

    for i, (dato, valor) in enumerate(info.items(), start=4):
        ws.cell(row=i - 1, column=1, value=dato).font = Font(bold=True, color=NEGRO)
        ws.cell(row=i - 1, column=2, value=valor).font = Font(color=GRIS)

    estilo_encabezado(ws, fila_tabla)
    ajustar_columnas(ws, desde_fila=fila_tabla)


def escribir_hoja(writer, nombre, df, color_pestana):
    from openpyxl.styles import Font, PatternFill

    if df.empty:
        df = pd.DataFrame({"sin datos": []})
    limpiar_df(df).to_excel(writer, sheet_name=nombre, index=False)
    ws = writer.sheets[nombre]
    ws.sheet_properties.tabColor = color_pestana
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    estilo_encabezado(ws, 1)

    columnas = {c.value: c.column for c in ws[1]}
    if "estado_descarga" in columnas:
        for (celda,) in ws.iter_rows(min_row=2, min_col=columnas["estado_descarga"], max_col=columnas["estado_descarga"]):
            if celda.value in COLOR_ESTADO:
                relleno, letra = COLOR_ESTADO[celda.value]
                celda.fill = PatternFill("solid", fgColor=relleno)
                celda.font = Font(color=letra, bold=True)
    if "url_web" in columnas and ws.max_row < 20000:
        for (celda,) in ws.iter_rows(min_row=2, min_col=columnas["url_web"], max_col=columnas["url_web"]):
            if celda.value:
                celda.hyperlink = celda.value
                celda.font = Font(color="4D6B12", underline="single")
    ajustar_columnas(ws)


def estilo_encabezado(ws, fila):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    for celda in ws[fila]:
        if celda.value is None:
            continue
        celda.fill = PatternFill("solid", fgColor=NEGRO)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.alignment = Alignment(vertical="center")
        celda.border = Border(bottom=Side(style="medium", color=VERDE))
    ws.row_dimensions[fila].height = 22


def limpiar_df(df):
    return df.apply(lambda col: col.map(lambda v: ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v))


def ajustar_columnas(ws, desde_fila=1):
    for columna in ws.iter_cols(min_row=desde_fila, max_row=min(ws.max_row, desde_fila + 200)):
        largo = max((len(str(c.value)) for c in columna if c.value is not None), default=8)
        ws.column_dimensions[columna[0].column_letter].width = min(max(largo + 2, 10), 60)


# ---------------------- #
#          MAIN          #
# ---------------------- #

def main():
    parser = argparse.ArgumentParser(description="Exporta el contenido de Tableau Server / Cloud a una carpeta local.")
    parser.add_argument("--config", default="config.ini", help="Archivo de configuración (por defecto config.ini)")
    parser.add_argument("--solo-inventario", action="store_true", help="Genera solo el inventario, sin descargar")
    parser.add_argument("--prueba", action="store_true",
                        help="Modo prueba: procesa solo unos pocos elementos por tipo (cantidad_prueba en config)")
    args = parser.parse_args()

    cfg = leer_config(args.config)
    t, e = cfg["tableau"], cfg["exportacion"]
    opciones = {
        "descargar_workbooks": e.getboolean("descargar_workbooks", True),
        "descargar_fuentes": e.getboolean("descargar_fuentes", True),
        "descargar_flujos": e.getboolean("descargar_flujos", True),
        "incluir_extractos": e.getboolean("incluir_extractos", True),
        "sobrescribir": e.getboolean("sobrescribir", False),
        "inventario_campos": e.getboolean("inventario_campos", True),
        "filtro_proyectos": [p.strip() for p in e.get("filtro_proyectos", "").split(",") if p.strip()],
        "modo_prueba": args.prueba or e.getboolean("modo_prueba", False),
        "cantidad_prueba": e.getint("cantidad_prueba", 5),
    }

    marca = datetime.now().strftime("%Y%m%d_%H%M%S")
    carpeta = os.path.abspath(e.get("carpeta_salida", "./exportacion"))
    ruta_log = configurar_logs(carpeta, marca)
    inicio = time.time()

    token_name = t.get("token_name") or os.environ.get("TABLEAU_TOKEN_NAME") or input("Nombre del token (PAT): ").strip()
    token_value = t.get("token_value") or os.environ.get("TABLEAU_TOKEN_VALUE") or getpass.getpass("Valor del token (PAT): ").strip()
    sitios = [s.strip() for s in t.get("sites", "").split(",")]

    modo = "Solo inventario" if args.solo_inventario else "Inventario + descarga"
    if opciones["modo_prueba"]:
        modo += f" | PRUEBA ({opciones['cantidad_prueba']} por tipo)"
    log.info("=" * 70)
    log.info(f"Exportador Tableau | {t['server_url']} | sitios: {', '.join(s or 'Default' for s in sitios)}")
    log.info(f"Modo: {modo}")
    log.info(f"Carpeta de salida: {carpeta}")
    log.info("=" * 70)

    conexion = Conexion(t["server_url"], token_name, token_value, t.getboolean("verify_ssl", True))
    sitios_info = []
    ruta_inventario = os.path.join(carpeta, f"inventario_{marca}.xlsx")
    ruta_pendientes = os.path.join(carpeta, f"pendientes_{marca}.xlsx")

    def info_ejecucion():
        minutos, segundos = divmod(int(time.time() - inicio), 60)
        return {
            "Servidor": t["server_url"], "Fecha de extracción": datetime.now().strftime("%Y-%m-%d %H:%M"),
            "Modo": modo, "Versión API": conexion.server.version, "Duración": f"{minutos:02d}:{segundos:02d}",
            "Incluye extractos": "Sí" if opciones["incluir_extractos"] else "No",
            "Elementos pendientes": len(pendientes),
        }

    try:
        if sitios == ["*"]:
            conexion.conectar("")
            sitios = [s.content_url for s in listar(conexion.server.sites)]
            log.info(f"Sitios encontrados: {', '.join(s or 'Default' for s in sitios)}")

        # ---- Fase 1: inventario de todos los sitios ----
        log.info("\n📋 FASE 1: INVENTARIO")
        for sitio in sitios:
            log.info(f"\n🔨 Inventariando sitio: {sitio or 'Default'}")
            try:
                conexion.conectar(sitio)
                info_sitio = leer_inventario(conexion, opciones)
                info_sitio["carpeta"] = os.path.join(carpeta, limpiar_nombre(sitio or "Default"))
                sitios_info.append(info_sitio)
            except Exception as ex:
                log.exception(f"   🚨 Error inventariando el sitio '{sitio or 'Default'}': {ex}")
                registrar_pendiente(sitio, "Sitio", "", sitio or "Default", "", t["server_url"], ex,
                                    "Revisar nombre del sitio, que el token pertenezca a ese sitio y re-ejecutar.")

        ruta_inventario = guardar_inventario(sitios_info, info_ejecucion(), ruta_inventario)
        log.info(f"\n📊 Inventario guardado: {ruta_inventario}")

        # ---- Fase 2: descarga ----
        if not args.solo_inventario:
            log.info("\n⬇  FASE 2: DESCARGA")
            for info_sitio in sitios_info:
                sitio = info_sitio["sitio"]
                log.info(f"\n🔨 Descargando sitio: {sitio or 'Default'}")
                try:
                    conexion.conectar(sitio)
                    descargas = descargar_sitio(conexion, info_sitio, opciones)
                    aplicar_descargas(info_sitio["tablas"], descargas, False)
                except Exception as ex:
                    log.exception(f"   🚨 Error descargando el sitio '{sitio or 'Default'}': {ex}")
                    registrar_pendiente(sitio, "Sitio", "", sitio or "Default", "", t["server_url"], ex,
                                        "Revisar el log y re-ejecutar; los archivos ya descargados se saltan.")
        else:
            for info_sitio in sitios_info:
                aplicar_descargas(info_sitio["tablas"], {}, True)

    except KeyboardInterrupt:
        log.warning("\n⛔ Ejecución interrumpida por el usuario. Se guarda lo obtenido hasta ahora.")
    finally:
        conexion.desconectar()
        info = info_ejecucion()
        try:
            ruta_inventario = guardar_inventario(sitios_info, info, ruta_inventario)
            ruta_pendientes = guardar_pendientes(info, ruta_pendientes)
        except Exception as ex:
            log.exception(f"🚨 No se pudieron guardar los Excel: {ex}")

        log.info("\n" + "=" * 70)
        for s in sitios_info:
            r = s["tablas"]["_resumen"]
            log.info(f"Sitio {r['sitio']}: {r['workbooks']} workbooks, {r['fuentes_de_datos']} fuentes, "
                     f"{r['flujos']} flujos | descargados {r['descargados']}, ya existían {r['ya_existian']}, "
                     f"con error {r['con_error']}")
        log.info(f"📊 Inventario : {ruta_inventario}")
        log.info(f"📝 Pendientes : {ruta_pendientes} ({len(pendientes)} elementos)")
        log.info(f"🗒  Log        : {ruta_log}")
        log.info(f"⏱  Duración   : {info['Duración']}")
        log.info("=" * 70)


if __name__ == "__main__":
    main()
